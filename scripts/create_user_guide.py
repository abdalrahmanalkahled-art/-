from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image


PROJECT_DIR = Path("/home/ubuntu/madar-marketing-manager")
ASSETS_DIR = PROJECT_DIR / "docs" / "user-guide-assets"
OUTPUT_PATH = PROJECT_DIR / "docs" / "دليل_استخدام_مدير_تسويق_مدار.docx"

PRIMARY = "1E5ED8"
SECONDARY = "0E9F6E"
TEXT = "1F2937"
MUTED = "5B6470"
LIGHT = "F2F6FF"
BORDER = "D7E2F4"
FONT_NAME = "Arial"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_rtl_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = alignment
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    return paragraph


def set_rtl_run(run, size=None, bold=None, color=None):
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT_NAME)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run._element.get_or_add_rPr().append(OxmlElement("w:rtl"))
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_text(document, text, size=11, bold=False, color=TEXT, space_after=6, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph = document.add_paragraph()
    set_rtl_paragraph(paragraph, alignment)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.35
    set_rtl_run(paragraph.add_run(text), size=size, bold=bold, color=color)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph()
    set_rtl_paragraph(paragraph)
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(6)
    sizes = {1: 18, 2: 15, 3: 12}
    colors = {1: PRIMARY, 2: SECONDARY, 3: TEXT}
    run = set_rtl_run(paragraph.add_run(text), size=sizes.get(level, 11), bold=True, color=colors.get(level, TEXT))
    if level == 1:
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "6")
        bottom.set(qn("w:color"), PRIMARY)
        borders.append(bottom)
        p_pr.append(borders)
    return paragraph


def add_bullet(document, text, level=0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_rtl_paragraph(paragraph)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.25
    set_rtl_run(paragraph.add_run(text), size=10.5, color=TEXT)
    return paragraph


def add_numbered_step(document, number, text):
    paragraph = document.add_paragraph()
    set_rtl_paragraph(paragraph)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.3
    set_rtl_run(paragraph.add_run(f"{number}. "), size=10.5, bold=True, color=PRIMARY)
    set_rtl_run(paragraph.add_run(text), size=10.5, color=TEXT)
    return paragraph


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    bidi.set(qn("w:val"), "1")
    tbl_pr.append(bidi)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_rtl(table)
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        if widths:
            cell.width = Cm(widths[index])
        set_cell_shading(cell, PRIMARY)
        set_cell_borders(cell, PRIMARY)
        paragraph = cell.paragraphs[0]
        set_rtl_paragraph(paragraph)
        paragraph.paragraph_format.space_after = Pt(0)
        set_rtl_run(paragraph.add_run(header), size=10, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if widths:
                cells[index].width = Cm(widths[index])
            if row_index % 2 == 0:
                set_cell_shading(cells[index], LIGHT)
            set_cell_borders(cells[index])
            paragraph = cells[index].paragraphs[0]
            set_rtl_paragraph(paragraph)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            set_rtl_run(paragraph.add_run(str(value)), size=9.5, color=TEXT)
    document.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_screenshot(document, filename, caption):
    image_path = ASSETS_DIR / filename
    if not image_path.exists():
        return
    compatible_path = image_path
    if image_path.suffix.lower() == ".webp":
        compatible_path = image_path.with_suffix(".png")
        if not compatible_path.exists():
            with Image.open(image_path) as source:
                source.convert("RGB").save(compatible_path, "PNG", optimize=True)
    document.add_picture(str(compatible_path), width=Inches(6.85))
    paragraph = document.paragraphs[-1]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph = document.add_paragraph()
    set_rtl_paragraph(caption_paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    caption_paragraph.paragraph_format.space_after = Pt(10)
    set_rtl_run(caption_paragraph.add_run(caption), size=9, color=MUTED)


def add_callout(document, title, text, color=PRIMARY):
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    set_table_rtl(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_borders(cell, color)
    paragraph = cell.paragraphs[0]
    set_rtl_paragraph(paragraph)
    paragraph.paragraph_format.space_after = Pt(0)
    set_rtl_run(paragraph.add_run(f"{title}: "), size=10, bold=True, color=color)
    set_rtl_run(paragraph.add_run(text), size=10, color=TEXT)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("صفحة ")
    set_rtl_run(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT_NAME)
    normal.font.size = Pt(11)
    for style_name in ("List Bullet", "List Bullet 2"):
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:cs"), FONT_NAME)
    header = section.header.paragraphs[0]
    set_rtl_paragraph(header)
    set_rtl_run(header.add_run("مدير تسويق مدار | دليل الاستخدام"), size=9, color=PRIMARY)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_cover(document):
    document.add_paragraph().paragraph_format.space_before = Pt(28)
    logo = PROJECT_DIR / "assets" / "images" / "icon.png"
    if logo.exists():
        document.add_picture(str(logo), width=Inches(1.35))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = document.add_paragraph()
    set_rtl_paragraph(title, WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(8)
    set_rtl_run(title.add_run("دليل الاستخدام التفصيلي"), size=26, bold=True, color=PRIMARY)
    subtitle = document.add_paragraph()
    set_rtl_paragraph(subtitle, WD_ALIGN_PARAGRAPH.CENTER)
    subtitle.paragraph_format.space_after = Pt(6)
    set_rtl_run(subtitle.add_run("تطبيق مدير تسويق مدار"), size=20, bold=True, color=TEXT)
    description = document.add_paragraph()
    set_rtl_paragraph(description, WD_ALIGN_PARAGRAPH.CENTER)
    description.paragraph_format.space_before = Pt(8)
    description.paragraph_format.space_after = Pt(16)
    description.paragraph_format.line_spacing = 1.35
    set_rtl_run(description.add_run("مرجع عملي لإدارة التسويق الميداني، المحلات، الاستبيانات، الفعاليات، المخزون، التقارير والتحليلات."), size=12, color=MUTED)
    add_callout(document, "الإصدار الموثق", "النسخة الحالية من التطبيق المتاحة ضمن المشروع، مع لقطات شاشة من معاينة الواجهات الأساسية.")
    add_text(document, "إعداد: فريق مدير تسويق مدار", size=10.5, color=MUTED, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(document, "تاريخ الإعداد: 16 آب 2026", size=10.5, color=MUTED, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    document.add_page_break()


def add_contents(document):
    add_heading(document, "فهرس الدليل", 1)
    contents = [
        "1. نبذة عن التطبيق ومجال الاستخدام",
        "2. البدء السريع والتنقل بين التبويبات",
        "3. لوحة التحكم الرئيسية",
        "4. إدارة المحلات",
        "5. الاستبيانات السوقية والدورات",
        "6. إدارة الفعاليات والتوثيق",
        "7. وحدات المزيد والإدارة المرجعية",
        "8. التحليلات المتقدمة وتقارير PDF وExcel",
        "9. إدارة الملفات والنسخ الاحتياطي والحماية",
        "10. إرشادات التشغيل وحل المشكلات الشائعة",
    ]
    for item in contents:
        add_bullet(document, item)
    add_callout(document, "ملاحظة", "تُخزّن بيانات التشغيل في النسخة الحالية محلياً على الجهاز. لذلك احرص على استخدام النسخ الاحتياطي بانتظام قبل تغيير الهاتف أو إزالة التطبيق.", SECONDARY)
    document.add_page_break()


def add_overview(document):
    add_heading(document, "1. نبذة عن التطبيق ومجال الاستخدام", 1)
    add_text(document, "مدير تسويق مدار هو تطبيق جوال لإدارة عمليات التسويق الميداني من نقطة واحدة. يساعد فريق التسويق على تنظيم بيانات المحلات، إنشاء استبيانات السوق وتنفيذها على شكل دورات، ربط المنتجات بالماركات والمناطق، إدارة الفعاليات واللوحات والستاندات، ومتابعة المخزون والصرفيات والأهداف التسويقية.")
    add_text(document, "يركّز التطبيق على العمل الميداني القابل للقياس. فعند تسجيل نتائج الاستبيانات أو الفعاليات، تُستخدم هذه البيانات لاحقاً في التحليلات والتقارير، مما يسهّل تتبع التواجد حسب المنتج أو المنطقة أو المحل أو دورة الاستبيان.")
    add_table(document, ["المجال", "ما الذي يتيحه التطبيق؟"], [
        ["المحلات", "حفظ بيانات المحل ومالكه ورقم الاتصال والمنطقة والتصنيف والملاحظات."],
        ["الاستبيانات", "إنشاء قوالب، ترتيب الأصناف والمنتجات، تنفيذ دورات، تسجيل التواجد والرفوف والأسعار والتعليقات."],
        ["الفعاليات", "تخطيط الفعاليات وربطها بالأهداف، وتسجيل موادها وتوثيقها الإعلامي."],
        ["المستودع", "إدارة المواد وحركتها والأدوات المرتبطة بالماركات وحالتها وصورها."],
        ["التحليلات", "مخططات ومرشحات وتقارير PDF وExcel عربية مبنية على البيانات المسجلة."],
    ], [4.1, 12.7])
    add_heading(document, "المتطلبات الأساسية", 2)
    add_bullet(document, "هاتف Android أو iPhone حديث يعمل بالإصدار المخصص للتطبيق.")
    add_bullet(document, "منح صلاحية الصور/الوسائط عند استخدام توثيق الفعاليات أو صورة المحل، بحسب طلب النظام.")
    add_bullet(document, "وجود بيانات فعلية، مثل محلات وقوالب ونتائج استبيانات، للاستفادة الكاملة من التحليلات والتقارير.")
    add_bullet(document, "يفضل عمل نسخة احتياطية للبيانات قبل حذف التطبيق أو الانتقال إلى هاتف جديد.")


def add_navigation(document):
    add_heading(document, "2. البدء السريع والتنقل بين التبويبات", 1)
    add_text(document, "بعد فتح التطبيق تظهر خمسة تبويبات ثابتة أسفل الشاشة. تنتقل بينها بلمسة واحدة؛ وتُعرض الصفحات بحركة انتقال خفيفة للمحافظة على سياق المستخدم.")
    add_table(document, ["التبويب", "الاستخدام الرئيسي"], [
        ["الرئيسية", "متابعة المؤشرات السريعة والانتقال إلى إجراءات ميدانية شائعة."],
        ["المحلات", "إضافة المحلات وتعديلها والبحث عنها وتصنيفها."],
        ["الاستبيانات", "إدارة القوالب والنتائج والتحليلات الخاصة بالاستبيانات."],
        ["الفعاليات", "إنشاء الفعاليات وتتبّع حالتها وتوثيقها."],
        ["المزيد", "الوحدات الإدارية المتخصصة، الإعدادات، الوضع الليلي وتسجيل الخروج."],
    ], [3.4, 13.4])
    add_heading(document, "مسار عمل مقترح للمستخدم الجديد", 2)
    for number, text in enumerate([
        "ابدأ بإضافة الماركات والمناطق من «المزيد»، ثم أضف منتجات الشركة والمنافسين واربطها بالماركات المناسبة.",
        "أضف المحلات وحدد المنطقة والتصنيف وبيانات التواصل لكي تصبح جاهزة للاستخدام في الاستبيانات والفعاليات.",
        "أنشئ قالب استبيان، ورتّب الأصناف والمنتجات كما تريد أن تظهر للمندوب أثناء الزيارة.",
        "نفذ الاستبيان على المحلات، وسجل التواجد والرفوف والسعر والتعليقات والصورة عند الحاجة.",
        "راجع التحليلات المتقدمة، ثم صدّر PDF أو Excel وشاركه عبر التطبيقات المتاحة على الهاتف.",
    ], start=1):
        add_numbered_step(document, number, text)


def add_dashboard(document):
    add_heading(document, "3. لوحة التحكم الرئيسية", 1)
    add_text(document, "لوحة التحكم هي نقطة المتابعة اليومية. تعرض بصورة مختصرة الميزانية المستخدمة، وعدد الفعاليات، وعدد المحلات النشطة، وعدد الاستبيانات، ومتوسط التواجد. وتحت هذه المؤشرات توجد إجراءات سريعة لتقليل عدد الخطوات اللازمة لإضافة فعالية أو استبيان أو حركة مخزون أو صرفية.")
    add_screenshot(document, "dashboard.webp", "الشكل 1: لوحة التحكم الرئيسية في معاينة التطبيق.")
    add_heading(document, "كيفية الاستفادة اليومية", 2)
    add_bullet(document, "راجع بطاقة الميزانية المستخدمة لمعرفة الوضع العام للمصروفات مقابل الميزانية المسجلة.")
    add_bullet(document, "استخدم «فعالية جديدة» عند التخطيط لنشاط ميداني، أو «استبيان جديد» للوصول السريع إلى دورة الدراسة السوقية.")
    add_bullet(document, "راقب مؤشري المحلات والتواجد لتتأكد من نمو قاعدة الزيارات وتحسن تغطية المنتجات.")


def add_stores(document):
    add_heading(document, "4. إدارة المحلات", 1)
    add_text(document, "يضم تبويب المحلات سجل نقاط البيع التي يتعامل معها الفريق. يدعم البحث بالاسم أو المنطقة، والتصنيف، والإضافة والتعديل والحذف مع تأكيد قبل الإجراء الحساس. تتكامل بيانات المحل مباشرة مع الاستبيانات والفعاليات والتحليلات، لذلك يُنصح بإدخال الاسم والمنطقة وتصنيف المحل بدقة.")
    add_screenshot(document, "stores.webp", "الشكل 2: شاشة إدارة المحلات والبحث والتصنيف والزر العائم للإضافة.")
    add_heading(document, "إضافة محل", 2)
    for number, text in enumerate([
        "افتح تبويب «المحلات» ثم اضغط زر الإضافة العائم.",
        "أدخل اسم المحل، وبيانات المالك أو المسؤول، ورقم التواصل والمنطقة والعنوان إن توفرت.",
        "اختر التصنيف المناسب من القائمة. يمكن للمستخدم إضافة تصنيفات جديدة أو تعديلها أو حذفها بحسب سياسة الفريق.",
        "احفظ النموذج. بعد الحفظ الناجح يبدأ نموذج الإضافة فارغاً حتى لا تتكرر بيانات المحل السابق.",
    ], start=1):
        add_numbered_step(document, number, text)
    add_callout(document, "تنبيه مهم", "عند حذف تصنيف محل، يُعاد التعامل مع المحلات المرتبطة به وفق الخيار البديل المتاح في نافذة الإدارة؛ لذلك راجع هذا الأثر قبل التأكيد.", "F59E0B")
    add_heading(document, "الممارسات الموصى بها", 2)
    add_bullet(document, "استخدم تسمية موحدة للمناطق حتى لا تنقسم النتائج التحليلية بين أسماء متقاربة لنفس المنطقة.")
    add_bullet(document, "أدخل رقم تواصل صالحاً لأن هذا الحقل يظهر في جدول المحلات المدروسة عند تصدير تقارير Excel.")


def add_surveys(document):
    add_heading(document, "5. الاستبيانات السوقية والدورات", 1)
    add_text(document, "الاستبيان في التطبيق يتكون من قالب ودورات تنفيذ. القالب يحدد الأصناف والمنتجات والأسئلة والخيارات، بينما الدورة تمثل فترة جمع بيانات مستقلة. هذا التنظيم يسمح بمقارنة التواجد عبر الزمن دون خلط النتائج القديمة والجديدة.")
    add_screenshot(document, "surveys.webp", "الشكل 3: تبويب الاستبيانات، ويضم القوالب والنتائج والتحليلات.")
    add_heading(document, "إنشاء قالب استبيان", 2)
    for number, text in enumerate([
        "من تبويب «قوالب» أنشئ استبياناً جديداً وحدد اسمه ووصفه عند الحاجة.",
        "أضف الأصناف ثم المنتجات. يحتفظ التطبيق بترتيب أول اختيار للأصناف والمنتجات ليظهر بالترتيب نفسه عند تنفيذ الاستبيان.",
        "فعّل «مع نسبة الظهور» إذا كان الاستبيان يحتاج قياس المساحة أو الرفوف، وفعّل «مع سعر المنتج» إذا كانت الأسعار مطلوبة.",
        "أضف الأسئلة الاختيارية وحدد ما إذا كانت الإجابة أحادية أو متعددة، ثم احفظ القالب.",
    ], start=1):
        add_numbered_step(document, number, text)
    add_heading(document, "تنفيذ استبيان على محل", 2)
    for number, text in enumerate([
        "اختر القالب والدورة الحالية، ثم اختر محلاً من القائمة الموحّدة للمحلات.",
        "لا يسمح التطبيق بأخذ الاستبيان للمحل نفسه أكثر من مرة داخل الدورة الواحدة؛ ويعرض تنبيهاً داخل نافذة اختيار المحل عند وجود نتيجة سابقة.",
        "عند تفعيل نسبة الظهور، أدخل إجمالي رفوف المحل مرة واحدة ثم أدخل عدد الرفوف المشغولة لكل منتج؛ يحتسب التطبيق النسبة تلقائياً.",
        "أدخل السعر عند تفعيل خياره، وسجل التعليقات مع تصنيفها، ويمكن التقاط صورة للمحل إذا كان هذا الخيار مفعلاً في القالب.",
        "احفظ النتيجة. ستظهر لاحقاً في «النتائج» وفي التحليلات والتقارير.",
    ], start=1):
        add_numbered_step(document, number, text)
    add_heading(document, "إنهاء الدورة", 2)
    add_text(document, "من بطاقة الاستبيان يمكن إنهاء الدورة الحالية. تبقى النتائج السابقة محفوظة وتبدأ الدورة التالية كفترة مستقلة، ليتمكن المستخدم من مقارنة الأداء بمرور الوقت. يحمل اسم الدورة المنتهية فترة زمنية مبنية على أول وآخر نتيجة ضمنها.")


def add_events(document):
    add_heading(document, "6. إدارة الفعاليات والتوثيق", 1)
    add_text(document, "يستخدم تبويب الفعاليات لتسجيل النشاطات التسويقية الميدانية ومتابعة حالتها. يدعم البحث وتصفية الفعاليات حسب الحالة: الكل، مخططة، جارية، مكتملة أو ملغاة. ويمكن ربط الفعالية بهدف من الخطة التسويقية، فتظهر مساهمتها في التقدم التحليلي للماركة أو الهدف ذي الصلة.")
    add_screenshot(document, "events.webp", "الشكل 4: شاشة إدارة الفعاليات مع مرشحات الحالة والبحث.")
    add_heading(document, "إضافة فعالية وتوثيقها", 2)
    for number, text in enumerate([
        "افتح تبويب «الفعاليات» واضغط زر الإضافة العائم.",
        "أدخل عنوان الفعالية وتاريخها وموقعها وحالتها، واختر الهدف المرتبط عند توفره.",
        "سجل الميزانية والمواد أو أي ملاحظات تشغيلية مطلوبة وفق نموذج الفعالية.",
        "أضف صوراً أو فيديوهات التوثيق عند الحاجة. يعرض التطبيق معاينة مصغرة للوسائط، ويمكن فتح الفيديو بمشغّل متاح على الهاتف أو مشاركته عبر ورقة المشاركة النظامية.",
        "احفظ الفعالية، ثم حدّث حالتها مع تقدم التنفيذ حتى تظهر بشكل صحيح في مؤشرات المتابعة والتحليلات.",
    ], start=1):
        add_numbered_step(document, number, text)
    add_callout(document, "نصيحة", "اربط الفعالية بهدف له ماركة محددة عندما يكون الهدف من النشاط دعم ماركة بعينها؛ فهذا يعطي التحليلات المتقدمة أساساً أوضح لعرض الإنجاز.", SECONDARY)


def add_more_modules(document):
    add_heading(document, "7. وحدات المزيد والإدارة المرجعية", 1)
    add_text(document, "تبويب «المزيد» هو بوابة الوحدات المتخصصة وإعدادات التطبيق. تُفتح هذه الوحدات في صفحة مستقلة داخل التطبيق، ويظهر زر إغلاق واضح للعودة إلى صفحة المزيد.")
    add_screenshot(document, "more.webp", "الشكل 5: تبويب المزيد، ويعرض الوحدات الإدارية والإعدادات في شبكة موحّدة.")
    add_table(document, ["الوحدة", "الاستخدام العملي"], [
        ["إدارة المستودع", "إضافة مواد المخزون وتعديلها وحذفها، مراجعة سجل الحركة، وإدارة الأدوات وحالتها وربطها بالماركات."],
        ["إدارة الصرفيات", "تسجيل المصروفات، اختيار تصنيف قابل للإدارة، وربط الصرفيات بسياق العمل عند الحاجة."],
        ["الخطة التسويقية", "إضافة الأهداف والمهام ومؤشرات الأداء، وربط الهدف بماركة اختيارية."],
        ["اللوحات والستاندات", "متابعة اللوحات الإعلانية والستاندات، المواقع والحالات والصور والارتباط بالماركات والمناطق."],
        ["التقارير", "الوصول إلى الملفات المحفوظة من عمليات التصدير وإدارتها مع التحديد المتعدد والحذف المؤكد."],
        ["إدارة المنتجات", "إدارة منتجات الشركة والمنافسين وتصنيفاتها وربط المنتجات بالماركات."],
        ["الماركات والمناطق", "إنشاء الماركات والمناطق وتقييم المناطق وتصنيفها لتصبح مرجعاً موحداً لبقية الوحدات."],
        ["التحليلات المتقدمة", "تجميع بيانات الاستبيانات والفعاليات واللوحات والستاندات وإعداد التقارير التحليلية."],
    ], [4.2, 12.6])
    add_heading(document, "إدارة المستودع والأدوات", 2)
    add_bullet(document, "عند إضافة مادة، اختر الفئة من قائمة قابلة للإدارة وحدد رمزها البصري، ثم سجّل الكمية ووحدة القياس المناسبة.")
    add_bullet(document, "عند حذف مادة، يحذف التطبيق سجل الحركة المرتبط بها أيضاً بعد نافذة تأكيد واضحة؛ راجع أثر الحذف قبل التنفيذ.")
    add_bullet(document, "في تبويب الأدوات، يُسجّل اسم الأداة وعدد القطع وحالتها وصورتها الاختيارية. ويمكن ربطها بماركة واحدة أو عدة ماركات بشرط ألا يتجاوز عدد الماركات المحددة عدد القطع.")
    add_heading(document, "الماركات والمناطق والمنتجات", 2)
    add_text(document, "الماركة هي المحور المشترك بين المنتجات والفعاليات واللوحات والستاندات والأهداف. أما المنطقة فهي مرجع موحّد للمحلات والاستبيانات واللوحات. إدخال هذين العنصرين أولاً يرفع دقة المرشحات والتحليلات اللاحقة.")
    add_heading(document, "الإعدادات", 2)
    add_bullet(document, "فعّل الوضع الليلي أو النهاري من المفتاح الموجود في قسم الإعدادات لتغيير المظهر دون تغيير البيانات.")
    add_bullet(document, "استخدم النسخ الاحتياطي بصورة دورية لحماية البيانات المحلية، ثم استخدم تسجيل الخروج عند إنهاء الجلسة على جهاز مشترك.")


def add_analytics_reports(document):
    add_heading(document, "8. التحليلات المتقدمة وتقارير PDF وExcel", 1)
    add_text(document, "تجمع التحليلات المتقدمة بيانات الاستبيانات والدورات والمنتجات والماركات والمناطق والمحلات والفعاليات واللوحات والستاندات. يعتمد المحتوى المعروض على المرشحات التي يحددها المستخدم؛ لذا تُعرض رسالة عدم كفاية البيانات عندما لا توجد نتائج مسجلة توافق الاختيارات الحالية.")
    add_heading(document, "تبويبات التحليل", 2)
    add_table(document, ["التبويب", "ما الذي يحلله؟"], [
        ["الاستبيانات والدورات", "نسب التواجد والظهور والسعر حسب الاستبيان والدورة، مع فصل الرسوم والجداول حسب الصنف."],
        ["تتبع محل أو منطقة", "اتجاه التواجد زمنياً لكل منتج في محل واحد أو جميع محلات المنطقة، مع التعليقات المصنفة."],
        ["الفعاليات واللوحات", "تقدم الأهداف المرتبطة بالماركات، مع تفصيل الفعاليات واللوحات والستاندات ضمن الماركات والمناطق المحددة."],
    ], [4.8, 12.0])
    add_heading(document, "استخدام التحليلات المتقدمة", 2)
    for number, text in enumerate([
        "افتح «المزيد» ثم «التحليلات المتقدمة».",
        "اختر الاستبيان أو الدورة أو الماركة أو المنطقة أو المحل بحسب نوع التبويب، ثم حدد المنتجات المطلوبة عند توفر هذا الخيار.",
        "راجع المؤشرات والمخططات والجداول. تُفصل الأصناف في المخططات حتى لا تختلط منتجات أصناف مختلفة في رسم واحد.",
        "استخدم زر «الإعدادات» من الزر العائم لاختيار نوع المخطط وإضافة شعار التقرير والتحكم في الجداول وتسميات المخطط.",
        "من الزر العائم نفسه اختر التصدير إلى PDF أو Excel لتوليد تقرير مطابق للمرشحات النشطة.",
    ], start=1):
        add_numbered_step(document, number, text)
    add_heading(document, "إعدادات المخططات", 2)
    add_bullet(document, "يمكن اختيار مخطط خطي أو أعمدة أو دائري أو مساحي أو نقاط أو رادار أو شمعدان. النوع المختار يُطبّق على التبويبات ويدخل في تقرير PDF.")
    add_bullet(document, "يمكن تفعيل عرض نسبة التواجد واسم المنتج على المخطط. ينظّم التطبيق توزيع الأعمدة والمسافات تلقائياً بحسب عدد العناصر لتقليل التداخل.")
    add_bullet(document, "يمكن إضافة شعار من الجهاز في رأس تقرير PDF، وتحديد الجداول التي يراد تضمينها، مثل جدول المحلات المدروسة.")
    add_heading(document, "محتوى التقرير", 2)
    add_table(document, ["الصيغة", "المحتوى وطريقة الاستخدام"], [
        ["PDF", "تقرير عربي من اليمين إلى اليسار يتضمن نطاق التقرير والمؤشرات والمخططات والجداول. تُفصل مخططات الأصناف ويظهر جدول المنتجات والدورات أسفل كل مخطط عند توفر البيانات."],
        ["Excel", "مصنف عربي منظم بأوراق مستقلة للملخص والمنتجات والدورات والتحليل التسويقي. ويُضاف جدول المحلات المدروسة في ورقة مستقلة عند اختياره، متضمناً الحالة لكل مادة والملاحظات."],
    ], [3.6, 13.2])
    add_callout(document, "آلية الحفظ والمشاركة", "بعد إنشاء التقرير يتحقق التطبيق من نجاح الملف ثم يسجله في مركز التقارير ويعرض خيارات المشاركة المتاحة على الهاتف. إذا تعذر حفظ ملف محلياً، استخدم خيار المشاركة لإرساله إلى تطبيق مناسب مثل البريد أو التخزين السحابي.", "F59E0B")


def add_safety_troubleshooting(document):
    add_heading(document, "9. إدارة الملفات والنسخ الاحتياطي والحماية", 1)
    add_text(document, "لأن التطبيق موجه للعمل الميداني، فإن حماية البيانات ودقة إدخالها أمران أساسيان. جميع عمليات الحذف الحساسة تستخدم نافذة تأكيد موحّدة؛ لا تؤكد الحذف قبل مراجعة تأثيره على السجلات المرتبطة، خصوصاً المواد التي يرتبط بها سجل حركة أو الماركات المرتبطة بمنتجات وفعاليات ولوحات وستاندات.")
    add_heading(document, "إرشادات الحماية", 2)
    add_bullet(document, "أنشئ نسخة احتياطية دورية قبل تحديث التطبيق أو تغيير الجهاز.")
    add_bullet(document, "استخدم تسجيل الخروج عندما يستخدم الجهاز أكثر من عضو من الفريق.")
    add_bullet(document, "احرص على تسمية موحدة للماركات والمناطق والمنتجات لتجنب نتائج تحليلية موزعة بين مسميات متشابهة.")
    add_bullet(document, "تحقق من الصور والفيديوهات بعد إرفاقها في الفعالية، ثم استخدم المشاركة لمراجعتها أو أرشفتها خارج الجهاز عند الحاجة.")
    add_heading(document, "10. إرشادات التشغيل وحل المشكلات الشائعة", 1)
    add_table(document, ["الحالة", "الإجراء المقترح"], [
        ["لا توجد بيانات كافية للتحليل", "أضف محلات وقالب استبيان ونتائج ضمن الدورة، ثم راجع أن مرشحات الماركة والمنطقة والمحل توافق هذه النتائج."],
        ["لا يمكن اختيار محل في الاستبيان", "تأكد من إضافة المحل أولاً، وتحقق مما إذا كان قد أُخذ له استبيان في الدورة الحالية؛ يمنع النظام التكرار ضمن الدورة."],
        ["التقرير لا يحوي الجدول المطلوب", "افتح إعدادات التحليلات وفعل جدول المحلات المدروسة أو الجداول الأخرى قبل إعادة التصدير."],
        ["محتوى النموذج اختفى خلف لوحة المفاتيح", "مرّر محتوى النافذة إلى الحقول السفلية. تم إعداد النوافذ الأساسية لتتحرك مع لوحة المفاتيح وتبقي أزرار الحفظ في موضع يمكن الوصول إليه."],
        ["لا يظهر زر الرجوع في نافذة داخلية", "استخدم زر الإغلاق أو العودة أعلى النافذة، أو إيماءة الرجوع في الهاتف. احفظ أو ألغِ أي مسودة قبل الخروج إذا ظهرت نافذة تأكيد."],
        ["تعذر بناء تطبيق جديد عبر الخدمة السحابية", "تحقق من رسالة خدمة البناء. إذا كانت متعلقة بحصة بناء أو اشتراك، فهي ليست خللاً في بيانات التطبيق أو تقريره وتحتاج معالجة من جهة خدمة البناء."],
    ], [5.0, 11.8])
    add_heading(document, "خلاصة عملية", 2)
    add_text(document, "للحصول على أفضل نتيجة، ابدأ من المراجع الأساسية: الماركات والمناطق والمنتجات، ثم المحلات، ثم قوالب الاستبيان والدورات. بعد جمع نتائج موثوقة ستتمكن من استخراج تحليل زمني واضح وتقرير قابل للمشاركة يدعم اتخاذ القرار الميداني.")


def add_reference_note(document):
    document.add_page_break()
    add_heading(document, "مصدر الدليل", 1)
    add_text(document, "أُعد هذا الدليل بالاعتماد على النسخة الحالية من واجهات ومكونات تطبيق مدير تسويق مدار، وعلى لقطات الشاشة الملتقطة من معاينة التطبيق أثناء إعداد الدليل. وهو يشرح سلوك الاستخدام المقصود للوظائف المتاحة في النسخة الموثقة.", size=10.5)
    add_callout(document, "ملاحظة للإصدار", "قد تختلف مواضع العناصر أو التفاصيل البصرية قليلاً بين شاشة المعاينة والهواتف الفعلية، بينما يبقى مسار الاستخدام والوظائف الأساسية كما هو موضح في هذا الدليل.")


def build_document():
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_cover(document)
    add_contents(document)
    add_overview(document)
    add_navigation(document)
    add_dashboard(document)
    add_stores(document)
    add_surveys(document)
    add_events(document)
    add_more_modules(document)
    add_analytics_reports(document)
    add_safety_troubleshooting(document)
    add_reference_note(document)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
